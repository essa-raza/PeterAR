from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def workbook_values(path: Path) -> dict[str, list[list[object]]]:
    workbook = load_workbook(path, data_only=True)
    values: dict[str, list[list[object]]] = {}
    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        rows: list[list[object]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        values[sheet] = rows
    return values


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare two workbook outputs by sheet names and cell values.")
    parser.add_argument("left", type=Path)
    parser.add_argument("right", type=Path)
    args = parser.parse_args()

    left = workbook_values(args.left)
    right = workbook_values(args.right)

    if set(left) != set(right):
        raise SystemExit(f"Sheet names differ: {sorted(left)} != {sorted(right)}")

    for sheet_name in left:
        if left[sheet_name] != right[sheet_name]:
            raise SystemExit(f"Workbook values differ in sheet '{sheet_name}'.")

    print("Workbook outputs match.")
