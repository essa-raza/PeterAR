from __future__ import annotations

import csv
import time
from collections import defaultdict, deque
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Deque

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


CSV_HEADERS = [
    "Type",
    "Omschrijving",
    "Datum",
    "Factuurnummer",
    "Vervaldatum",
    "Openstaand",
    "Bedrag",
    "Match status",
]
OUTPUT_HEADERS = CSV_HEADERS[:7] + ["Commentaar", "Mail", "Whatsapp"]
LEGACY_NOTE_COLUMNS_WITH_STATUS = (9, 10, 11)
LEGACY_NOTE_COLUMNS_NO_STATUS = (8, 9, 10)


@dataclass
class LegacyRow:
    signature: tuple[str, ...]
    notes: tuple[object, object, object]
    row_fill: object | None


@dataclass
class MergeResult:
    output_path: Path
    matched_count: int
    unmatched_customers: list[str]
    total_customer_sections: int
    debug_log: list[str]
    log_file_path: Path | None = None


class MergeLogger:
    def __init__(
        self,
        ui_callback: Callable[[str], None] | None = None,
        save_detailed_log: bool = False,
    ) -> None:
        self.ui_callback = ui_callback
        self.save_detailed_log = save_detailed_log
        self.lines: list[str] = []
        self._started_at = time.perf_counter()

    def log(self, message: str, ui_message: str | None = None, include_in_file: bool = True) -> None:
        elapsed = time.perf_counter() - self._started_at
        line = f"[{elapsed:7.2f}s] {message}"
        if include_in_file and self.save_detailed_log:
            self.lines.append(line)
        if self.ui_callback and ui_message:
            self.ui_callback(ui_message)

    def log_path(self, label: str, path: Path) -> None:
        resolved = path.expanduser().resolve(strict=False)
        self.log(f"{label}: {resolved}", include_in_file=True)

    def log_step_duration(self, label: str, started_at: float) -> None:
        self.log(f"{label} completed in {time.perf_counter() - started_at:.2f}s")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def row_signature(values: list[object], width: int = 8) -> tuple[str, ...]:
    padded = list(values[:width])
    if len(padded) < width:
        padded.extend([""] * (width - len(padded)))
    return tuple(clean_text(value) for value in padded)


def parse_csv_sections(csv_path: Path, logger: MergeLogger | None = None) -> tuple[list[str], list[list[str]]]:
    started_at = time.perf_counter()
    if logger:
        logger.log("Opening customer CSV for reading.", ui_message="Opening CSV file...")
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(handle, delimiter=delimiter)
        rows = list(reader)
    if logger:
        logger.log(f"Detected '{delimiter}' as the CSV delimiter.")
        logger.log_step_duration("Reading customer CSV", started_at)

    if not rows:
        raise ValueError("CSV file is empty.")

    header = list(rows[0][:8])
    if header != CSV_HEADERS:
        raise ValueError(
            f"Unexpected CSV headers. Expected {CSV_HEADERS} but received {header}."
        )

    data_rows: list[list[str]] = []
    for row in rows[1:]:
        trimmed = list(row[:8])
        if len(trimmed) < 8:
            trimmed.extend([""] * (8 - len(trimmed)))
        data_rows.append(trimmed)

    return header, data_rows


def build_legacy_row_map(
    xlsx_path: Path,
    logger: MergeLogger | None = None,
) -> tuple[dict[tuple[str, ...], Deque[LegacyRow]], int]:
    started_at = time.perf_counter()
    if logger:
        logger.log("Opening source workbook for reading.", ui_message="Opening source workbook...")
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    header_values = [sheet.cell(1, col_idx).value for col_idx in range(1, sheet.max_column + 1)]
    normalized_headers = [clean_text(value).lower() for value in header_values]
    if normalized_headers[:11] == [value.lower() for value in CSV_HEADERS + ["Commentaar", "Mail", "Whatsapp"]]:
        note_columns = LEGACY_NOTE_COLUMNS_WITH_STATUS
        match_width = 8
    elif normalized_headers[:10] == [value.lower() for value in OUTPUT_HEADERS]:
        note_columns = LEGACY_NOTE_COLUMNS_NO_STATUS
        match_width = 7
    else:
        note_columns = LEGACY_NOTE_COLUMNS_WITH_STATUS
        match_width = 8

    if logger:
        logger.log(
            f"Workbook matching mode uses {match_width} signature columns and note columns {note_columns}."
        )

    row_map: dict[tuple[str, ...], Deque[LegacyRow]] = defaultdict(deque)
    for row_idx in range(2, sheet.max_row + 1):
        signature = row_signature(
            [sheet.cell(row_idx, col_idx).value for col_idx in range(1, match_width + 1)],
            width=match_width,
        )
        notes = tuple(sheet.cell(row_idx, col_idx).value for col_idx in note_columns)

        row_fill = None
        for col_idx in range(1, min(sheet.max_column, 11) + 1):
            fill = sheet.cell(row_idx, col_idx).fill
            if fill and fill.fill_type:
                row_fill = copy(fill)
                break

        row_map[signature].append(LegacyRow(signature=signature, notes=notes, row_fill=row_fill))

    if logger:
        logger.log(
            f"Indexed {sheet.max_row - 1} legacy workbook rows for matching.",
            ui_message="Reading workbook rows...",
        )
        logger.log_step_duration("Reading source workbook", started_at)
    return row_map, match_width


def default_output_path(csv_path: Path) -> Path:
    return csv_path.with_name(f"{csv_path.stem} - merged.xlsx")


def format_output_sheet(sheet) -> None:
    widths = {
        "A": 22,
        "B": 90,
        "C": 14,
        "D": 18,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 45,
        "I": 18,
        "J": 18,
    }
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = wrap_alignment

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"


def write_output(
    output_path: Path,
    header: list[str],
    data_rows: list[list[str]],
    legacy_rows: dict[tuple[str, ...], Deque[LegacyRow]],
    match_width: int,
    logger: MergeLogger | None = None,
) -> tuple[int, list[str], int]:
    started_at = time.perf_counter()
    if logger:
        logger.log("Creating output workbook in memory.", ui_message="Building merged workbook...")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merged"
    sheet.append(OUTPUT_HEADERS)

    matched = 0
    unmatched: list[str] = []

    for row_idx, row in enumerate(data_rows, start=2):
        padded = list(row)[: len(header)]
        if len(padded) < len(header):
            padded.extend([""] * (len(header) - len(padded)))

        legacy_row = None
        signature = row_signature(padded, width=match_width)
        matches = legacy_rows.get(signature)
        if matches:
            legacy_row = matches.popleft()
            matched += 1
        else:
            unmatched.append(" | ".join(signature))

        notes = list(legacy_row.notes) if legacy_row else ["", "", ""]
        output_row = padded[:7] + notes
        sheet.append(output_row)

        if legacy_row and legacy_row.row_fill:
            for col_idx in range(1, 11):
                sheet.cell(row_idx, col_idx).fill = copy(legacy_row.row_fill)

    if logger:
        logger.log(
            f"Prepared {len(data_rows)} output rows with {matched} matched rows and {len(unmatched)} unmatched rows.",
            ui_message="Merging notes into output rows...",
        )
        logger.log("Applying workbook formatting.", ui_message="Formatting workbook...")
    format_output_sheet(sheet)
    if logger:
        logger.log("Saving merged workbook to disk.", ui_message="Saving output workbook...")
    workbook.save(output_path)
    if logger:
        logger.log_step_duration("Writing output workbook", started_at)
    return matched, unmatched, len(data_rows)


def write_debug_log(log_file_path: Path, lines: list[str]) -> None:
    log_file_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def merge_customer_notes(
    csv_path: Path | str,
    excel_path: Path | str,
    output_path: Path | str | None = None,
    log_callback: Callable[[str], None] | None = None,
    save_log_file: bool = False,
) -> MergeResult:
    overall_started_at = time.perf_counter()
    csv_file = Path(csv_path)
    excel_file = Path(excel_path)
    output_file = Path(output_path) if output_path else default_output_path(csv_file)
    logger = MergeLogger(ui_callback=log_callback, save_detailed_log=save_log_file)

    logger.log("Starting merge process.", ui_message="Starting merge...")
    logger.log_path("Customer CSV path", csv_file)
    logger.log_path("Source workbook path", excel_file)
    logger.log_path("Output workbook path", output_file)

    header, data_rows = parse_csv_sections(csv_file, logger=logger)
    legacy_rows, match_width = build_legacy_row_map(excel_file, logger=logger)
    matched, unmatched, total = write_output(
        output_file, header, data_rows, legacy_rows, match_width, logger=logger
    )

    logger.log(
        f"Merge finished in {time.perf_counter() - overall_started_at:.2f}s.",
        ui_message="Merge completed successfully.",
    )

    log_file_path: Path | None = None
    if save_log_file:
        log_file_path = output_file.with_suffix(".log")
        logger.log_path("Log file path", log_file_path)
        write_debug_log(log_file_path, logger.lines)

    return MergeResult(
        output_path=output_file,
        matched_count=matched,
        unmatched_customers=unmatched,
        total_customer_sections=total,
        debug_log=logger.lines,
        log_file_path=log_file_path,
    )
